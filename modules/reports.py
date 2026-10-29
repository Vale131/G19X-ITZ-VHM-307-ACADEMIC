import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import io

def generar_reporte_excel(df_historico, df_forecast, df_cartera, res_mc):
    """Genera un libro de Excel profesional consolidado en memoria con openpyxl."""
    wb = Workbook()
    
    # Estilos Corporativos PluriOne
    header_fill = PatternFill(start_color="6D1B2A", end_color="6D1B2A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    # 1. Hoja de Resumen Ejecutivo y Riesgo VaR
    ws_resumen = wb.active
    ws_resumen.title = "Resumen & VaR"
    ws_resumen.views.sheetView[0].showGridLines = True
    
    ws_resumen.append(["Métrica de Riesgo / Indicador", "Valor Calculado"])
    ws_resumen.append(["Value at Risk 95% (VaR 95%) ($)", res_mc.get('var_95', 0)])
    ws_resumen.append(["Value at Risk 99% (VaR 99%) ($)", res_mc.get('var_99', 0)])
    ws_resumen.append(["Probabilidad de Déficit de Caja (%)", res_mc.get('prob_deficit', 0)])
    ws_resumen.append(["Saldo Final Esperado ($)", res_mc.get('saldo_final_esperado', 0)])

    for row in ws_resumen.iter_rows(min_row=1, max_row=5, min_col=1, max_col=2):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin

    for cell in ws_resumen[1]:
        cell.fill = header_fill
        cell.font = header_font

    # 2. Hoja de Proyección de Flujo
    ws_proj = wb.create_sheet(title="Proyección de Flujo")
    ws_proj.views.sheetView[0].showGridLines = True
    if df_forecast is not None and not df_forecast.empty:
        # Soporte dinámico para columnas de proyección
        cols_deseadas = ['ds', 'Ingresos', 'Egresos', 'Flujo_Neto', 'Saldo_Acumulado', 'flujo_proyectado', 'saldo_proyectado']
        cols_existentes = [c for c in cols_deseadas if c in df_forecast.columns]
        
        df_export_proj = df_forecast[cols_existentes] if cols_existentes else df_forecast.copy()
        
        for r in dataframe_to_rows(df_export_proj, index=False, header=True):
            ws_proj.append(r)

        for cell in ws_proj[1]:
            cell.fill = header_fill
            cell.font = header_font

    # 3. Hoja de Cartera CxC / CxP
    ws_cart = wb.create_sheet(title="Cartera CxC y CxP")
    ws_cart.views.sheetView[0].showGridLines = True
    if df_cartera is not None and not df_cartera.empty:
        for r in dataframe_to_rows(df_cartera, index=False, header=True):
            ws_cart.append(r)

        for cell in ws_cart[1]:
            cell.fill = header_fill
            cell.font = header_font

    # 4. Hoja de Datos Históricos
    if df_historico is not None and not df_historico.empty:
        ws_hist = wb.create_sheet(title="Histórico de Caja")
        ws_hist.views.sheetView[0].showGridLines = True
        for r in dataframe_to_rows(df_historico, index=False, header=True):
            ws_hist.append(r)

        for cell in ws_hist[1]:
            cell.fill = header_fill
            cell.font = header_font

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()